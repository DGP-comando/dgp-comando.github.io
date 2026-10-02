#!/usr/bin/env python3
"""Gera as camadas das Unidades de Referência (URs) dos programas do IDR-Paraná.

Fonte: pasta "000-Ações Integradas" (RECURSOS NATURAIS E SUSTENTABILIDADE),
uma subpasta por programa:
  - Grãos/UIRTs_Programa_Graos.shp (EPSG:4326, DBF em UTF-8): 201 UIRTs com
    agricultor, técnicos responsáveis e MIP/MID/coletor de esporos. 4 linhas
    sem geometria: a coordenada sai das colunas Latitude/Longitude.
  - Café/Assistidos_Cafe.xls (xlrd): 62 propriedades, "lat, lon" do Google Maps.
  - Piscicultura/Unidades_Programa Estadual Piscicultura.xlsx: só a coordenada
    ("lat º, lon º"); o município sai do ponto no polígono do IBGE.
  - Pecuária de Corte/Anexos.zip: Unidades_ProgramasEstaduais.xlsx (Programa
    Purunã, Pecuária Moderna, Associação Purunã) e "Regional de Umuarama_4
    URs.kml", que é uma mensagem MIME: a 1ª parte é o KML das 4 URs de corte de
    Umuarama; a 2ª, em base64, é o mesmo Ortigueira_<produtor>.kml solto na
    pasta (polígono do imóvel no CAR). As 5 coordenadas preenchem as linhas da
    planilha que diziam "Enviado por whatsapp formato kml" (mesmo produtor).
  - Turismo Rural/shp/usodosolo_rotadoqueijo.shp (EPSG:31982): uso do solo
    dos imóveis do CAR das queijarias da Rota do Queijo. Preferido à camada
    usodosolo_rotadoqueijo do 00-VETORES/vetores.gpkg (116 feições, EPSG:4674):
    o gpkg é o shp deduplicado por OBJECTID, o que apagou 1 imóvel de
    Pinhão (área sobreposta a outro imóvel) e 4 classes de 1 imóvel de Toledo.
    No shp, as classes de cada imóvel somam a área
    do imóvel no CAR. Geometria idêntica em dois imóveis vira uma feição só,
    com os dois códigos.
  - Fora daqui: Mapa Rota do Queijo Paranaense.kml (27 queijarias, todas a
    menos de 100 m de um ponto da camada Rotas turísticas) e
    Agroindústrias/coordenadas.xlsx (1.023 linhas idênticas às do diagnóstico
    já em agroindustrias-idr-pr.geojson).

Coordenada: ler_coord() do build_agroindustrias_idr.py (formatos mistos,
lat/lon trocados), conferida contra o município declarado; o resultado fica em
"Checagem da coordenada". Sem coordenada utilizável (vazia ou fora do PR), o
ponto vai para o centro do município declarado e a checagem diz isso.

LGPD: nome do produtor + coordenada da propriedade (e código do imóvel no CAR)
= dado pessoal. Saída só em data/privado/ (bucket datageo-privado):
urs-graos-pr, urs-cafe-pr, urs-piscicultura-pr, urs-pecuaria-corte-pr e
usodosolo-queijarias-pr (.geojson, lon/lat SIRGAS 2000 ~ WGS84, 5 casas, UTF-8).

Uso:
  py -3 scripts/build_urs_programas.py
  SUPABASE_SERVICE_ROLE_KEY=... py -3 scripts/upload_privado.py urs- usodosolo-
"""

import base64
import json
import re
import zipfile
from email import message_from_bytes
from pathlib import Path
from xml.etree import ElementTree as ET

import geopandas as gpd
import openpyxl
import pandas as pd
from shapely.geometry import Point, mapping, shape
from shapely.ops import unary_union

from build_agroindustrias_idr import chave, ler_coord
from build_faxinais import Municipios
from build_limites_ambientais import _round_coords, _text

ROOT = Path(__file__).resolve().parent.parent
SRC = Path(r'H:\IDR-PARANA\RECURSOS NATURAIS E SUSTENTABILIDADE\000-Ações Integradas')
PRIV = ROOT / 'data' / 'privado'
KML = {'k': 'http://www.opengis.net/kml/2.2'}

OK = 'no município declarado'
PELO_PONTO = 'município pelo ponto (planilha sem município)'
# Grafias da origem que não batem com o IBGE nem sem acento.
ALIAS = {'dr ulysses': 'doutor ulysses'}
# "Joäo", "Peräo": trema no lugar do til na planilha.
TREMA = str.maketrans('äöÄÖ', 'ãõÃÕ')


class Base:
    """Municípios do IBGE por nome e o contorno do PR para ler_coord()."""

    def __init__(self):
        self.ibge = Municipios()
        self.geom = {cod: g for (cod, _), g in self.ibge.feicoes}
        self.nome = {cod: nm for (cod, nm), _ in self.ibge.feicoes}
        self.por_nome = {chave(nm): (cod, nm) for (cod, nm), _ in self.ibge.feicoes}
        self.pr = unary_union(list(self.geom.values())).buffer(0.01)

    def municipio(self, nome):
        """(código, nome oficial) ou None. "Califórnia (Ortigueira)" -> Ortigueira."""
        m = re.search(r'\(([^)]+)\)', nome or '')
        k = re.sub(r' doeste$', ' d oeste', chave(m.group(1) if m else nome))  # "Pérola DOeste"
        k = ALIAS.get(k, k)
        # "DIAMANTE DO OESTE" (ADAPAR) = "Diamante D'Oeste" (IBGE); só na falta do nome exato.
        return self.por_nome.get(k) or self.por_nome.get(re.sub(r' do oeste$', ' d oeste', k))

    def ponto(self, lat, lon, nome_mun):
        """(lon, lat, props de município e checagem) para uma linha da planilha.

        Com município declarado, confere a coordenada contra ele; sem município
        (piscicultura), o ponto define o município.
        """
        mun = self.municipio(nome_mun) if nome_mun else None
        if nome_mun and not mun:
            raise ValueError(f'município sem correspondência no IBGE: {nome_mun!r}')
        lat, lon = _text(lat), _text(lon)
        if generica(lat, lon):
            motivo = f'coordenada genérica ({lat}, {lon})'
            got = None
        else:
            motivo = f'coordenada fora do PR ({lat}, {lon})' if lat and lon else 'sem coordenada na planilha'
            got = lat and lon and ler_coord(lat, lon, self.geom[mun[0]] if mun else None, self.pr)
        if got and not mun:
            achado = self.ibge.em(Point(got[0], got[1]))
            return got[0], got[1], {'Município': achado[1], 'ibge': achado[0], 'Checagem da coordenada': PELO_PONTO}
        if got:
            return got[0], got[1], {'Município': mun[1], 'ibge': mun[0], 'Checagem da coordenada': got[2]}
        if not mun:
            return None
        c = self.geom[mun[0]].representative_point()
        return c.x, c.y, {'Município': mun[1], 'ibge': mun[0],
                          'Checagem da coordenada': f'{motivo}: ponto no centro do município'}


def generica(lat, lon):
    """"-24", "-52": graus inteiros nos dois eixos marcam preenchimento, não o lugar."""
    try:
        return float(lat).is_integer() and float(lon).is_integer()
    except ValueError:
        return False


def feature(lon, lat, props):
    props = {k: v for k, v in props.items() if v not in (None, '')}
    return {'type': 'Feature', 'properties': props,
            'geometry': {'type': 'Point', 'coordinates': [round(lon, 5), round(lat, 5)]}}


def latlon(txt):
    """'-23.53º, -51.39º' -> ('-23.53', '-51.39'); outro texto -> ('', '')."""
    partes = [p.strip() for p in str(txt or '').replace('º', '').split(',')]
    return (partes[0], partes[1]) if len(partes) == 2 and all(re.fullmatch(r'-?\d+\.\d+', p) for p in partes) else ('', '')


def write(name, features):
    out = PRIV / name
    out.write_text(json.dumps({'type': 'FeatureCollection', 'features': features},
                              ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    falhas = sum(1 for f in features if f['properties'].get('Checagem da coordenada', OK) not in (OK, PELO_PONTO))
    print(f'{name}: {len(features)} feições, {falhas} não conferem com o município, {out.stat().st_size / 1024:.0f} KB')


# ------------------------------------------------------------------- grãos

def build_graos(base):
    g = gpd.read_file(SRC / 'Grãos' / 'UIRTs_Programa_Graos.shp', encoding='utf-8')
    feats = []
    for r in g.itertuples(index=False):
        lat, lon = (r.geometry.y, r.geometry.x) if r.geometry is not None else (r.Latitude, r.Longitude)
        x, y, mun = base.ponto(lat, lon, r.Municipio)
        feats.append(feature(x, y, {
            'Produtor': _text(r.Agricultor), **mun, 'Regional': _text(r.Região),
            'Responsáveis': _text(r.Responsáv).strip('[]'),
            'MIP': _text(r.MIP), 'MID': _text(r.MID), 'Coletor de esporos': _text(r.Coletor),
        }))
    write('urs-graos-pr.geojson', dedup(feats))


def dedup(feats):
    """Mesmo produtor no mesmo ponto: uma UR só (responsáveis somados). Mesmo
    produtor no mesmo município com uma linha conferida: as que não conferem saem."""
    por_ponto = {}
    for f in feats:
        p = f['properties']
        k = (chave(p.get('Produtor')), tuple(f['geometry']['coordinates']))
        if k in por_ponto:
            q = por_ponto[k]['properties']
            resp = [s.strip() for s in f"{q.get('Responsáveis', '')}, {p.get('Responsáveis', '')}".split(',')]
            q['Responsáveis'] = ', '.join(dict.fromkeys(s for s in resp if s))
            continue
        por_ponto[k] = f
    unicos = list(por_ponto.values())
    conferidos = {(chave(f['properties'].get('Produtor')), f['properties']['ibge']) for f in unicos
                  if f['properties']['Checagem da coordenada'] == OK}
    saida = [f for f in unicos if f['properties']['Checagem da coordenada'] == OK
             or (chave(f['properties'].get('Produtor')), f['properties']['ibge']) not in conferidos]
    if len(saida) < len(feats):
        print(f'  {len(feats) - len(saida)} linhas repetidas fundidas')
    return saida


# -------------------------------------------------------------------- café

def build_cafe(base):
    df = pd.read_excel(SRC / 'Café' / 'Assistidos_Cafe.xls', dtype=str)
    feats = []
    for r in df.itertuples(index=False):
        lat, lon = latlon(r[2])
        x, y, mun = base.ponto(lat, lon, _text(r[4]))
        feats.append(feature(x, y, {'Produtor': _text(r[3]), **mun}))
    write('urs-cafe-pr.geojson', feats)


# ------------------------------------------------------------ piscicultura

def build_piscicultura(base):
    ws = openpyxl.load_workbook(SRC / 'Piscicultura' / 'Unidades_Programa Estadual Piscicultura.xlsx',
                                read_only=True, data_only=True).worksheets[0]
    feats, vazias = [], 0
    for n, _, coord in list(ws.iter_rows(values_only=True))[1:]:
        lat, lon = latlon(coord)
        got = base.ponto(lat, lon, None) if lat else None
        if not got:
            vazias += 1
            continue
        x, y, mun = got
        feats.append(feature(x, y, {'Unidade': f'UR {n}', **mun}))
    print(f'  piscicultura: {vazias} linhas numeradas sem coordenada (planilha só tem {len(feats)} unidades)')
    write('urs-piscicultura-pr.geojson', feats)


# -------------------------------------------------------- pecuária de corte

def kml_placemarks(xml):
    """[(nome, geometria shapely, {SimpleData})] de um KML."""
    out = []
    for pm in ET.fromstring(xml).iter(f'{{{KML["k"]}}}Placemark'):
        nome = _text(pm.findtext('k:name', '', KML))
        dados = {sd.get('name'): _text(sd.text) for sd in pm.iter(f'{{{KML["k"]}}}SimpleData')}
        pt = pm.find('.//k:Point/k:coordinates', KML)
        if pt is not None:
            lon, lat = map(float, pt.text.strip().split(',')[:2])
            out.append((nome, Point(lon, lat), dados))
            continue
        ring = pm.find('.//k:outerBoundaryIs//k:coordinates', KML)
        if ring is not None:
            pts = [tuple(map(float, c.split(',')[:2])) for c in ring.text.split()]
            out.append((nome, shape({'type': 'Polygon', 'coordinates': [pts]}), dados))
    return out


def kmls_pecuaria():
    """{chave do produtor: (lon, lat, extra)} das URs com coordenada só em KML."""
    pasta = SRC / 'Pecuária de Corte'
    with zipfile.ZipFile(pasta / 'Anexos.zip') as z:
        bruto = z.read('Regional de Umuarama_4 URs.kml')
    kml1, resto = bruto.split(b'</kml>', 1)
    # Nome do produtor sai do nome do arquivo (Ortigueira_<produtor>.kml), não do código.
    arq_ortigueira = next(pasta.glob('Ortigueira_*.kml'))
    produtor_ortigueira = arq_ortigueira.stem.split('_', 1)[1].strip()
    ortigueira = arq_ortigueira.read_bytes()
    # 2ª parte MIME (base64) = o KML de Ortigueira solto na pasta?
    anexo = message_from_bytes(b'Content-Type: multipart/mixed; boundary="' + resto.split(b'\n', 2)[1][2:].strip()
                               + b'"\n\n' + resto.lstrip())
    partes = [p.get_payload(decode=True) for p in anexo.walk() if p.get_filename()]
    igual = any(p and p.strip() == ortigueira.strip() for p in partes)
    print(f'  KML de Umuarama: 2ª parte MIME {"idêntica" if igual else "DIFERENTE"} ao {arq_ortigueira.name}')

    achados = {}
    for nome, geom, _ in kml_placemarks(kml1 + b'</kml>'):
        produtor = nome.split('_')[0].strip()
        achados[chave(produtor)] = (geom.x, geom.y, {'Fonte da coordenada': 'KML URs_Corte_Umuarama (WhatsApp)',
                                                     'nome no KML': nome})
    for _, geom, d in kml_placemarks(ortigueira):
        c = geom.representative_point()
        achados[chave(produtor_ortigueira)] = (c.x, c.y, {
            'Fonte da coordenada': 'polígono do imóvel no CAR (KML), ponto no centro',
            'Imóvel CAR': d.get('recibo'),
            'Área do imóvel (ha)': round(float(d['area']), 2) if d.get('area') else None,
            'Módulos fiscais': round(float(d['modfiscais']), 2) if d.get('modfiscais') else None,
        })
    return achados


def build_pecuaria(base):
    with zipfile.ZipFile(SRC / 'Pecuária de Corte' / 'Anexos.zip') as z:
        ws = openpyxl.load_workbook(z.open('Unidades_ProgramasEstaduais.xlsx'), data_only=True).worksheets[0]
    kml = kmls_pecuaria()
    feats, sem_mun, usados = [], [], set()
    for row in list(ws.iter_rows(values_only=True))[1:]:
        _, programa, municipio, regional, produtor, coord, obs = (row + (None,) * 7)[:7]
        if not programa:
            continue
        produtor = _text(produtor).translate(TREMA)
        if not _text(municipio):
            sem_mun.append((programa, regional, _text(obs)))
            continue
        lat, lon = latlon(coord)
        extra = {}
        k = chave(produtor)
        if not lat and k in kml:
            lon, lat, extra = kml[k]
            usados.add(k)
        x, y, mun = base.ponto(lat, lon, _text(municipio))
        feats.append(feature(x, y, {
            'Produtor': produtor, 'Programa': _text(programa), **mun, 'Regional': _text(regional),
            'Observação': _text(obs), **extra,
        }))
    for k in set(kml) - usados:  # UR do KML sem linha na planilha: entra sozinha
        lon, lat, extra = kml[k]
        x, y, mun = base.ponto(lat, lon, extra['nome no KML'].split('_')[-1])
        # As URs de corte de Umuarama são da Pecuária Moderna na planilha.
        feats.append(feature(x, y, {'Produtor': extra['nome no KML'].split('_')[0].strip(),
                                    'Programa': 'Pecuária Moderna', **mun, **extra,
                                    'Observação': 'só no KML da Regional de Umuarama, sem linha na planilha'}))
    for f in feats:
        f['properties'].pop('nome no KML', None)
    print(f'  pecuária: {len(usados)} coordenadas do KML casadas com a planilha, {len(kml) - len(usados)} novas; '
          f'{len(sem_mun)} linhas sem município fora: {sem_mun}')
    write('urs-pecuaria-corte-pr.geojson', feats)


# ----------------------------------------------------- uso do solo (queijo)

def build_usosolo(base):
    g = gpd.read_file(SRC / 'Turismo Rural' / 'shp' / 'usodosolo_rotadoqueijo.shp').to_crs(4674)
    grupos = {}
    for r in g.itertuples(index=False):
        k = (int(r.OBJECTID), r.geometry.wkb)
        grupos.setdefault(k, []).append(r)
    feats = []
    for rows in grupos.values():
        r = rows[0]
        ibge = r.cod_imovel.split('-')[1]
        props = {
            'Classe': _text(r.NIVEL_II), 'Nível I': _text(r.NIVEL_I), 'Nível III': _text(r.NIVEL_III),
            'Área (ha)': round(float(r.Area_ha), 2), 'Município': base.nome[ibge], 'ibge': ibge,
            'Imóvel CAR': ' · '.join(x.cod_imovel for x in rows),
            'Área do imóvel (ha)': round(float(r.num_area), 2), 'Módulos fiscais': round(float(r.mod_fiscal), 2),
            'Condição no CAR': _text(r.des_condic),
        }
        geom = mapping(r.geometry)
        feats.append({'type': 'Feature', 'properties': props,
                      'geometry': {'type': geom['type'], 'coordinates': _round_coords(geom['coordinates'])}})
    print(f'  uso do solo: {len(g)} feições do shp, {len(g) - len(feats)} geometrias repetidas entre imóveis fundidas, '
          f'{g.cod_imovel.nunique()} imóveis do CAR')
    write('usodosolo-queijarias-pr.geojson', feats)


def main():
    PRIV.mkdir(parents=True, exist_ok=True)
    base = Base()
    build_graos(base)
    build_cafe(base)
    build_piscicultura(base)
    build_pecuaria(base)
    build_usosolo(base)


if __name__ == '__main__':
    main()
